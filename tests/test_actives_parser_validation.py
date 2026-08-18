from datetime import date, datetime

from controllers.actives_parser import ActivesParserController
from models import Actives, Consignment, CounterActive, Location, MileageStart, MileageTrain, Relocate, Storage
from tests.conftest import make_active, make_design_number

controller = ActivesParserController(owner=None)


async def test_design_number_old_column_name(db_session):
    await make_active(db_session, "ES1040010328")
    dn_id = await make_design_number(db_session, "A2V00002691454")

    rows = [{"Актив": "ES1040010328", "Новая Позиция ТМЦ": "A2V00002691454"}]
    errors, valid_rows = await controller._validate_design_number(db_session, rows)

    assert errors == []
    assert valid_rows == [("ES1040010328", dn_id, "A2V00002691454")]


async def test_design_number_new_column_name(db_session):
    await make_active(db_session, "ES1040010328")
    dn_id = await make_design_number(db_session, "A2V00002691454")

    rows = [{"Актив": "ES1040010328", "Новый ТМЦ номер": "A2V00002691454"}]
    errors, valid_rows = await controller._validate_design_number(db_session, rows)

    assert errors == []
    assert valid_rows == [("ES1040010328", dn_id, "A2V00002691454")]


async def test_design_number_short_column_name(db_session):
    await make_active(db_session, "ES1040010328")
    dn_id = await make_design_number(db_session, "A2V00002691454")

    rows = [{"Актив": "ES1040010328", "Позиция ТМЦ": "A2V00002691454"}]
    errors, valid_rows = await controller._validate_design_number(db_session, rows)

    assert errors == []
    assert valid_rows == [("ES1040010328", dn_id, "A2V00002691454")]


async def test_design_number_column_missing(db_session):
    rows = [{"Актив": "ES1040010328", "Другая колонка": "x"}]
    errors, valid_rows = await controller._validate_design_number(db_session, rows)

    assert valid_rows == []
    assert len(errors) == 1
    assert "не найдена колонка" in errors[0]["message"]


async def test_design_number_new_column_empty_value(db_session):
    await make_active(db_session, "ES1040010328")

    rows = [{"Актив": "ES1040010328", "Новый ТМЦ номер": ""}]
    errors, valid_rows = await controller._validate_design_number(db_session, rows)

    assert valid_rows == []
    assert len(errors) == 1
    assert "пустое" in errors[0]["message"]


async def make_recount_active(
    db_session,
    active_number: str = "UL0000001",
    id_unit_type: int | None = None,
    with_mileage_start: bool = True,
    with_counter: bool = True,
) -> int:
    active = Actives(active_number=active_number, id_unit_type=id_unit_type)
    db_session.add(active)
    await db_session.flush()
    if with_mileage_start:
        db_session.add(MileageStart(id_active=active.id, milage=0, milage_const=100))
    if with_counter:
        db_session.add(CounterActive(
            id_active=active.id, id_counter_type=3, date=datetime(2023, 1, 1, 12, 0), value=0,
        ))
    await db_session.flush()
    return active.id


async def test_recount_mileage_column_missing(db_session):
    errors, valid_rows = await controller._validate_recount_mileage(db_session, [{"Другое": "x"}])

    assert valid_rows == []
    assert len(errors) == 1
    assert "не найдена колонка 'Актив'" in errors[0]["message"]


async def test_recount_mileage_active_not_found(db_session):
    errors, valid_rows = await controller._validate_recount_mileage(db_session, [{"Актив": "UL0000001"}])

    assert valid_rows == []
    assert len(errors) == 1
    assert "Актив не найден" in errors[0]["message"]


async def test_recount_mileage_missing_mileage_start(db_session):
    await make_recount_active(db_session, with_mileage_start=False)

    errors, valid_rows = await controller._validate_recount_mileage(db_session, [{"Актив": "UL0000001"}])

    assert valid_rows == []
    assert "mileage_start не найдена" in errors[0]["message"]


async def test_recount_mileage_missing_counter(db_session):
    await make_recount_active(db_session, with_counter=False)

    errors, valid_rows = await controller._validate_recount_mileage(db_session, [{"Актив": "UL0000001"}])

    assert valid_rows == []
    assert "Счётчик пробега" in errors[0]["message"]


async def test_recount_mileage_train_needs_no_counter(db_session):
    await make_recount_active(db_session, id_unit_type=1, with_counter=False)

    errors, valid_rows = await controller._validate_recount_mileage(db_session, [{"Актив": "UL0000001"}])

    assert errors == []
    assert len(valid_rows) == 1
    assert valid_rows[0]["is_train"] is True
    assert valid_rows[0]["total"] == 0


async def test_recount_mileage_total_from_relocate_history(db_session):
    active_id = await make_recount_active(db_session)

    loc_storage = Location(id_storage=1)
    loc_train = Location(id_train=5)
    db_session.add_all([loc_storage, loc_train])
    await db_session.flush()

    db_session.add_all([
        # storage -> train before the cutoff: the train's mileage is added
        Relocate(id_active=active_id, id_location_old=loc_storage.id,
                 id_location_new=loc_train.id, date=datetime(2022, 1, 10, 8, 0)),
        # a move after the cutoff is ignored
        Relocate(id_active=active_id, id_location_old=loc_train.id,
                 id_location_new=loc_storage.id, date=datetime(2023, 1, 1, 8, 0)),
        MileageTrain(id_train=5, mileage_average=10, date_average=date(2022, 2, 1)),
        MileageTrain(id_train=5, mileage_average=15, date_average=date(2022, 3, 1)),
        # after the cutoff — not summed
        MileageTrain(id_train=5, mileage_average=99, date_average=date(2022, 6, 1)),
        # a different train — not summed
        MileageTrain(id_train=7, mileage_average=50, date_average=date(2022, 2, 1)),
    ])
    await db_session.flush()

    errors, valid_rows = await controller._validate_recount_mileage(db_session, [{"Актив": "UL0000001"}])

    assert errors == []
    assert len(valid_rows) == 1
    assert valid_rows[0]["total"] == 25
    assert valid_rows[0]["is_train"] is False


def test_recount_mileage_sql_body_order_and_trains():
    valid_rows = [
        {"row_num": 1, "active_number": "UL0000001", "id_active": 10, "total": 25, "is_train": False,
         "milage_const": None, "insert_mileage_start": False},
        {"row_num": 2, "active_number": "ES0000001", "id_active": 20, "total": -5, "is_train": True,
         "milage_const": None, "insert_mileage_start": False},
    ]
    lines = ActivesParserController._build_recount_mileage_sql_body(valid_rows)

    assert len(lines) == 3
    assert "UPDATE public.mileage_start SET milage = COALESCE(milage_const, 0) + (25)" in lines[0]
    assert "WHERE id_active = 10" in lines[0]
    assert "UPDATE public.mileage_start" in lines[1] and "(-5)" in lines[1]
    # the counter is recomputed only for non-trains, and only after mileage_start
    assert "UPDATE public.counter_active" in lines[2]
    assert "function_get_mileage" in lines[2]
    assert "id_active = 10" in lines[2]


async def test_recount_mileage_const_column_update_and_insert(db_session):
    await make_recount_active(db_session, "UL0000001")
    # an asset without mileage_start but with milage_const set is not an error — INSERT follows
    await make_recount_active(db_session, "UL0000002", with_mileage_start=False)

    rows = [
        {"Актив": "UL0000001", "milage_const": 500},
        {"Актив": "UL0000002", "milage_const": "700.0"},
    ]
    errors, valid_rows = await controller._validate_recount_mileage(db_session, rows)

    assert errors == []
    assert valid_rows[0]["milage_const"] == 500
    assert valid_rows[0]["insert_mileage_start"] is False
    assert valid_rows[1]["milage_const"] == 700
    assert valid_rows[1]["insert_mileage_start"] is True


async def test_recount_mileage_const_zero_and_empty_ignored(db_session):
    await make_recount_active(db_session, "UL0000001")
    await make_recount_active(db_session, "UL0000002")

    rows = [
        {"Актив": "UL0000001", "milage_const": 0},
        {"Актив": "UL0000002", "milage_const": None},
    ]
    errors, valid_rows = await controller._validate_recount_mileage(db_session, rows)

    assert errors == []
    assert all(vr["milage_const"] is None for vr in valid_rows)


async def test_recount_mileage_const_header_with_nbsp_and_e_spelling(db_session):
    # a real file: the header 'mileage_const\xa0' — spelled with 'e' and a non-breaking space
    await make_recount_active(db_session, "UL0000001")

    rows = [{"Актив": "UL0000001", "mileage_const\xa0": 295544}]
    errors, valid_rows = await controller._validate_recount_mileage(db_session, rows)

    assert errors == []
    assert valid_rows[0]["milage_const"] == 295544


async def test_recount_mileage_const_invalid_value(db_session):
    await make_recount_active(db_session, "UL0000001")

    rows = [{"Актив": "UL0000001", "milage_const": "abc"}]
    errors, valid_rows = await controller._validate_recount_mileage(db_session, rows)

    assert valid_rows == []
    assert "Некорректное значение milage_const" in errors[0]["message"]


async def test_recount_mileage_const_missing_mileage_start_without_const(db_session):
    # without milage_const a missing mileage_start is still an error
    await make_recount_active(db_session, "UL0000001", with_mileage_start=False)

    rows = [{"Актив": "UL0000001", "milage_const": ""}]
    errors, valid_rows = await controller._validate_recount_mileage(db_session, rows)

    assert valid_rows == []
    assert "mileage_start не найдена" in errors[0]["message"]


def test_recount_mileage_sql_body_const_first():
    valid_rows = [
        {"row_num": 1, "active_number": "UL0000001", "id_active": 10, "total": 25, "is_train": False,
         "milage_const": 500, "insert_mileage_start": False},
        {"row_num": 2, "active_number": "UL0000002", "id_active": 20, "total": 0, "is_train": False,
         "milage_const": 700, "insert_mileage_start": True},
    ]
    lines = ActivesParserController._build_recount_mileage_sql_body(valid_rows)

    assert len(lines) == 6
    # the milage_const block comes first: UPDATE for an existing row, INSERT for a new one
    assert "SET milage_const = 500" in lines[0] and "WHERE id_active = 10" in lines[0]
    assert lines[1].startswith("INSERT INTO public.mileage_start")
    assert "VALUES (20, 700, 700, true)" in lines[1]
    # then the milage and counter recount
    assert "SET milage = COALESCE(milage_const, 0) + (25)" in lines[2]
    assert "SET milage = COALESCE(milage_const, 0) + (0)" in lines[3]
    assert "UPDATE public.counter_active" in lines[4]
    assert "UPDATE public.counter_active" in lines[5]


async def make_named_actives_refs(db_session) -> None:
    db_session.add_all([
        Storage(name="Основной склад, МСК"),
        Consignment(name="ЧСП Исправные"),
    ])
    await db_session.flush()
    await make_design_number(db_session, "DU0012929")


NAMED_ROW = {"Актив": "SPD452390", "ТМЦ": "DU0012929",
             "Положение": "Основной склад, МСК", "Партия": "ЧСП Исправные"}


async def test_create_named_actives_columns_missing(db_session):
    errors, valid_rows = await controller._validate_create_named_actives(
        db_session, [{"Актив": "SPD452390", "Другое": "x"}])

    assert valid_rows == []
    assert len(errors) == 1
    assert "не найдены колонки" in errors[0]["message"]
    assert "ТМЦ" in errors[0]["message"]


async def test_create_named_actives_happy_path(db_session):
    await make_named_actives_refs(db_session)

    errors, valid_rows = await controller._validate_create_named_actives(db_session, [dict(NAMED_ROW)])

    assert errors == []
    assert len(valid_rows) == 1
    vr = valid_rows[0]
    assert vr["active_number"] == "SPD452390"
    assert all(vr[k] for k in ("id_design_number", "id_storage", "id_consignment"))


async def test_create_named_actives_already_exists(db_session):
    await make_named_actives_refs(db_session)
    await make_active(db_session, "SPD452390")

    errors, valid_rows = await controller._validate_create_named_actives(db_session, [dict(NAMED_ROW)])

    assert valid_rows == []
    assert "уже существует" in errors[0]["message"]


async def test_create_named_actives_duplicate_in_file(db_session):
    await make_named_actives_refs(db_session)

    errors, valid_rows = await controller._validate_create_named_actives(
        db_session, [dict(NAMED_ROW), dict(NAMED_ROW)])

    assert len(valid_rows) == 1
    assert "Дубликат внутри файла" in errors[0]["message"]


async def test_create_named_actives_refs_not_found(db_session):
    await make_named_actives_refs(db_session)

    rows = [
        dict(NAMED_ROW, **{"ТМЦ": "NOPE"}),
        dict(NAMED_ROW, **{"Актив": "SPD452391", "Положение": "Нет такого"}),
        dict(NAMED_ROW, **{"Актив": "SPD452392", "Партия": "Нет такой"}),
    ]
    errors, valid_rows = await controller._validate_create_named_actives(db_session, rows)

    assert valid_rows == []
    messages = " | ".join(e["message"] for e in errors)
    assert "ТМЦ не найдена" in messages
    assert "Склад не найден" in messages
    assert "Партия не найдена" in messages


def test_create_named_actives_sql_body():
    valid_rows = [
        {"row_num": 1, "active_number": "SPD452390", "id_design_number": 7,
         "id_storage": 3, "id_consignment": 4},
        {"row_num": 2, "active_number": "SPD452391", "id_design_number": 8,
         "id_storage": 3, "id_consignment": 4},
    ]
    lines = ActivesParserController._build_create_named_actives_sql_body(valid_rows)
    sql = "\n".join(lines)

    assert sql.startswith("DO $$")
    assert sql.count("INSERT INTO public.location") == 2
    assert sql.count("INSERT INTO public.actives") == 2
    assert "'SPD452390'" in sql and "'SPD452391'" in sql
    # the lcn is handed out from storage.last_lcn under a lock and written back
    assert "SELECT last_lcn INTO lcn_3 FROM public.storage WHERE id = 3 FOR UPDATE;" in sql
    assert "UPDATE public.storage SET last_lcn = lcn_3 WHERE id = 3;" in sql
    # the asset number counter is unused — the numbers come from the file
    assert "iterator_number_last" not in sql


async def test_recount_mileage_active_number_column_name(db_session):
    await make_recount_active(db_session, "UL0000001")

    rows = [{"active_number": "UL0000001"}]
    errors, valid_rows = await controller._validate_recount_mileage(db_session, rows)

    assert errors == []
    assert valid_rows[0]["active_number"] == "UL0000001"


async def test_delete_actives_column_missing(db_session):
    errors, valid_rows = await controller._validate_delete_actives(db_session, [{"Другое": "x"}])

    assert valid_rows == []
    assert "не найдена колонка 'Актив'" in errors[0]["message"]


async def test_delete_actives_not_found_and_duplicate(db_session):
    await make_active(db_session, "SPD1077356")

    rows = [{"Актив": "SPD1077356"}, {"Актив": "SPD1077356"}, {"Актив": "NOPE"}]
    errors, valid_rows = await controller._validate_delete_actives(db_session, rows)

    assert len(valid_rows) == 1
    messages = " | ".join(e["message"] for e in errors)
    assert "Дубликат внутри файла" in messages
    assert "Актив не найден" in messages


async def make_order_dependency_tables(db_session) -> None:
    """Create minimal versions of the order dependency tables in SQLite.

    The production tables from ORDERS_DEPENDENCY_CHECKS are not needed as ORM
    models — validation reaches them with raw SELECTs, so single-column tables
    are enough for the tests. Tables already described in models.py (relocate,
    order_to_actives) are skipped.
    """
    from sqlalchemy import text as sql_text
    from controllers.actives_parser import ORDERS_DEPENDENCY_CHECKS
    from models import Base
    existing = set(Base.metadata.tables)
    columns: dict[str, set[str]] = {}
    for tbl, col in ORDERS_DEPENDENCY_CHECKS:
        if f"public.{tbl}" not in existing:
            columns.setdefault(tbl, set()).add(col)
    for tbl, cols in columns.items():
        await db_session.execute(sql_text(
            f"CREATE TABLE IF NOT EXISTS public.{tbl} ({', '.join(f'{c} integer' for c in sorted(cols))})"
        ))


async def test_delete_actives_blocked_by_dependencies(db_session):
    from models import Orders, Ptoir
    await make_order_dependency_tables(db_session)
    ok_id = await make_active(db_session, "SPD1077356")
    blocked_relocate = await make_active(db_session, "SPD1077357")
    db_session.add_all([
        Relocate(id_active=blocked_relocate, date=datetime(2023, 1, 1)),
        # the counter (created by a trigger on every asset), the maintenance record and
        # an empty order against it do NOT block the deletion
        CounterActive(id_active=ok_id, id_counter_type=3, date=datetime(2023, 1, 1), value=0),
    ])
    ptoir_ok = Ptoir(number_ptoir="ТО1", id_active=ok_id)
    db_session.add(ptoir_ok)
    await db_session.flush()
    db_session.add(Orders(order_number="З-1", id_ptoir=ptoir_ok.id))
    await db_session.flush()

    rows = [{"Актив": "SPD1077356"}, {"Актив": "SPD1077357"}]
    errors, valid_rows = await controller._validate_delete_actives(db_session, rows)

    assert [vr["id_active"] for vr in valid_rows] == [ok_id]
    messages = " | ".join(e["message"] for e in errors)
    assert "'SPD1077357'" in messages and "relocate" in messages


async def test_delete_actives_blocked_by_order_with_dependencies(db_session):
    from sqlalchemy import text as sql_text
    from models import Orders
    await make_order_dependency_tables(db_session)
    blocked_id = await make_active(db_session, "SPD1077358")
    order = Orders(order_number="З-2", id_active=blocked_id)
    db_session.add(order)
    await db_session.flush()
    # the order has labour costs — the asset is blocked
    await db_session.execute(sql_text(
        f"INSERT INTO public.labor_costs (id_order) VALUES ({order.id})"))

    errors, valid_rows = await controller._validate_delete_actives(
        db_session, [{"Актив": "SPD1077358"}])

    assert valid_rows == []
    assert "заказов актива 'SPD1077358'" in errors[0]["message"]
    assert "labor_costs" in errors[0]["message"]


async def test_delete_actives_active_number_column(db_session):
    await make_active(db_session, "SPD1077356")

    errors, valid_rows = await controller._validate_delete_actives(
        db_session, [{"active_number": "SPD1077356"}])

    assert errors == []
    assert valid_rows[0]["active_number"] == "SPD1077356"


def test_delete_actives_sql_body():
    valid_rows = [
        {"row_num": 1, "active_number": "SPD1077356", "id_active": 10, "id_location": 100},
        {"row_num": 2, "active_number": "SPD1077357", "id_active": 20, "id_location": None},
    ]
    lines = ActivesParserController._build_delete_actives_sql_body(valid_rows)
    sql = "\n".join(lines)

    # framing: the DBA trigger tr_abort_delete is disabled for the transaction
    assert lines[0] == "ALTER TABLE public.counter_active DISABLE TRIGGER tr_abort_delete;"
    assert lines[-1] == "ALTER TABLE public.counter_active ENABLE TRIGGER tr_abort_delete;"
    # an asset with a location: orders + ptoir_level_warning + ptoir + counter_active +
    # mileage_start + actives + a guarded location; without a location, six statements
    assert len(lines) == 15
    assert "DELETE FROM public.orders WHERE id_active = 10 OR id_ptoir IN" in lines[1]
    assert "DELETE FROM public.ptoir_level_warning WHERE id_ptoir IN" in lines[2]
    assert "DELETE FROM public.ptoir WHERE id_active = 10;" in lines[3]
    assert "DELETE FROM public.counter_active WHERE id_active = 10;" in lines[4]
    assert "DELETE FROM public.mileage_start WHERE id_active = 10;" in lines[5]
    assert "DELETE FROM public.actives WHERE id = 10;" in lines[6]
    assert "DELETE FROM public.location l WHERE l.id = 100" in lines[7]
    assert "NOT EXISTS" in lines[7] and "materials" in lines[7] and "relocate" in lines[7]
    assert sql.count("DELETE FROM public.actives") == 2
    assert sql.count("DELETE FROM public.location") == 1


async def make_create_actives_refs(db_session) -> None:
    from models import IteratorNumberLast, Storage as StorageModel, Consignment as ConsignmentModel
    db_session.add_all([
        IteratorNumberLast(id=2, number=100, description="Номер следующего актива"),
        StorageModel(name="Основной склад, МСК", last_lcn=10),
        ConsignmentModel(name="ЧСП Исправные"),
    ])
    await db_session.flush()
    await make_design_number(db_session, "DU0012929")


CREATE_ACTIVES_ROW_TAIL = {"Количество": 1, "Склад": "Основной склад, МСК",
                           "Тип актива": "SPD", "Партия": "ЧСП Исправные"}


async def test_create_actives_old_tmc_column_name(db_session):
    await make_create_actives_refs(db_session)

    rows = [{"Номер ТМЦ (DU,KP,A2V)": "DU0012929", **CREATE_ACTIVES_ROW_TAIL}]
    errors, valid_rows = await controller._validate_create_actives_rows(db_session, rows)

    assert errors == []
    assert len(valid_rows) == 1


async def test_create_actives_new_articul_column_name(db_session):
    await make_create_actives_refs(db_session)

    rows = [{"АРТИКУЛ": "DU0012929", **CREATE_ACTIVES_ROW_TAIL}]
    errors, valid_rows = await controller._validate_create_actives_rows(db_session, rows)

    assert errors == []
    assert len(valid_rows) == 1


async def test_create_actives_tmc_column_missing(db_session):
    await make_create_actives_refs(db_session)

    rows = [{"Другое": "x", **CREATE_ACTIVES_ROW_TAIL}]
    errors, valid_rows = await controller._validate_create_actives_rows(db_session, rows)

    assert valid_rows == []
    assert "не найдена колонка 'АРТИКУЛ'" in errors[0]["message"]


def test_reconstruct_created_active_numbers():
    valid_rows = [{"type_active": "SPV"}, {"type_active": "SPV"}, {"type_active": "SPD"}]
    numbers = ActivesParserController._reconstruct_created_active_numbers(valid_rows, counter_after=103)

    # ACTIVE_NUMBER_LENGTH = 10: the prefix plus zero-padded digits
    assert numbers == ["SPV0000101", "SPV0000102", "SPD0000103"]


def test_build_created_actives_xlsx():
    import base64
    from io import BytesIO
    import openpyxl

    b64 = ActivesParserController._build_created_actives_xlsx(
        [("SPV0000101", "251200001"), ("SPV0000102", None)])
    wb = openpyxl.load_workbook(BytesIO(base64.b64decode(b64)))
    rows = list(wb.active.iter_rows(values_only=True))

    assert rows[0] == ("active_number", "serial_number")
    assert rows[1] == ("SPV0000101", "251200001")
    assert rows[2] == ("SPV0000102", None)
