from io import BytesIO

import pytest
from openpyxl import load_workbook

from controllers.reports import (
    DESIGN_NUMBERS_HEADERS, DESIGN_NUMBERS_SQL, TECHCARDS_OPERATIONS_HEADERS, TECHCARDS_OPERATIONS_SQL, report_xlsx,
)


@pytest.mark.parametrize("sql, headers", [
    (TECHCARDS_OPERATIONS_SQL, TECHCARDS_OPERATIONS_HEADERS),
    (DESIGN_NUMBERS_SQL, DESIGN_NUMBERS_HEADERS),
])
async def test_report_opens_with_headers(pg_session, sql, headers):
    content = await report_xlsx(pg_session, sql, headers)

    sheet = load_workbook(BytesIO(content)).active
    assert [cell.value for cell in sheet[1]] == headers
    assert sheet.max_column == len(headers)
    assert sheet.freeze_panes == "A2"
