from datetime import datetime
from io import BytesIO

from litestar import Controller, get
from litestar.response import Response, Template
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_COLUMN_WIDTH = 60

TECHCARDS_OPERATIONS_SQL = """
select
  tech.name as tech, tech.description, typ.name as work_type,
  st_typ.status, link_op.task, link_op.task_name, link_op.working_hours,
  early.name as tech_in, early.description as description_in
from techcard tech
left join work_type typ on typ.id = tech.id_work_type
left join techcard_status st on st.id = tech.id_techcard_status
left join techcard_status_types st_typ on st_typ.id = st.id_techcard_status_types
left join techcard_work_operation link_op on link_op.id_techcard_head = tech.id
left join techcard early on early.id = link_op.id_techcard_early
order by tech.name, link_op.task
"""

TECHCARDS_OPERATIONS_HEADERS = [
    "Техкарта", "Описание", "Вид работ", "Статус", "Операция", "Наименование операции",
    "Трудоёмкость, ч", "Вложенная техкарта", "Описание вложенной техкарты",
]

DESIGN_NUMBERS_SQL = """
select
  des.number, des.description, unit.name as unit_type, des.is_serial,
  des.is_created_1c, gr.name as commodity_group, cod.name as commodity_code,
  st.status as design_number_status, meas.name as measure_type,
  des.is_repairable, pg.name as product_group, des.description_1c,
  des.main_number, c_gr.name as counter_group,
  string_agg(m_us.name, ',') as maint_use,
  des.is_serial_1c
from design_number des
left join unit_type unit on unit.id = des.id_unit_type
left join design_number_status st on st.id = des.id_design_number_status
left join product_group pg on pg.id = des.id_product_group
left join measure_type meas on meas.id = des.id_measure_type
left join commodity_group gr on gr.id = des.id_commodity_group
left join commodity_code cod on cod.id = des.id_commodity_code
left join counter_group c_gr on c_gr.id = des.id_counter_group
left join design_number_to_maint_use link on link.id_design_number = des.id
left join maint_use m_us on m_us.id = link.id_maint_use
group by des.number, des.description, unit_type, des.is_serial,
  des.is_created_1c, commodity_group, commodity_code,
  design_number_status, measure_type,
  des.is_repairable, product_group, des.description_1c,
  des.main_number, counter_group, des.is_serial_1c
order by des.number
"""

DESIGN_NUMBERS_HEADERS = [
    "Номер", "Описание", "Единица измерения", "Серийный", "Создан в 1С", "Товарная группа",
    "Товарный код", "Статус", "Тип измерения", "Ремонтопригодный", "Группа продукции",
    "Описание 1С", "Основной номер", "Группа счётчиков", "Применение в ТО", "Серийный в 1С",
]


def build_xlsx(headers: list[str], rows: list[tuple]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append([str(value).lower() if isinstance(value, bool) else value for value in row])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(sheet.iter_cols(values_only=True), start=1):
        width = max(len(str(value)) for value in column if value is not None)
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, MAX_COLUMN_WIDTH)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def report_xlsx(db_session: AsyncSession, sql: str, headers: list[str]) -> bytes:
    rows = (await db_session.execute(text(sql))).all()
    return build_xlsx(headers, [tuple(row) for row in rows])


def xlsx_response(content: bytes, name: str) -> Response:
    filename = f"{name}_{datetime.now():%Y-%m-%d_%H%M}.xlsx"
    return Response(
        content=content,
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


class ReportsController(Controller):
    path = "/reports"

    @get("/")
    async def index(self) -> Template:
        return Template(template_name="reports.html", context={"active_page": "reports"})

    @get("/techcards-operations")
    async def techcards_operations(self, db_session: AsyncSession) -> Response:
        content = await report_xlsx(db_session, TECHCARDS_OPERATIONS_SQL, TECHCARDS_OPERATIONS_HEADERS)
        return xlsx_response(content, "techcards_operations")

    @get("/design-numbers")
    async def design_numbers(self, db_session: AsyncSession) -> Response:
        content = await report_xlsx(db_session, DESIGN_NUMBERS_SQL, DESIGN_NUMBERS_HEADERS)
        return xlsx_response(content, "design_numbers")
